import csv
from pathlib import Path

from openpyxl import load_workbook

from csv_cleaner.cli import main

SAMPLE = (
    "Customer ID;E-mail;Signup Date\nC1;A@Example.com;01.02.2026\nC2;a@example.com;2026-02-01\n"
)


def _read(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def test_cli_writes_clean_csv_and_report(tmp_path, capsys):
    src = tmp_path / "in.csv"
    src.write_text(SAMPLE, encoding="utf-8")
    out = tmp_path / "out.csv"

    code = main(
        [
            str(src),
            "-o",
            str(out),
            "--dates",
            "signup_date",
            "--emails",
            "e_mail",
            "--key",
            "e_mail",
        ]
    )

    assert code == 0
    assert _read(out) == [
        {"customer_id": "C1", "e_mail": "a@example.com", "signup_date": "2026-02-01"}
    ]
    issues = _read(tmp_path / "out_issues.csv")
    assert issues[0]["problem"] == "duplicate removed"
    assert "rows out:    1" in capsys.readouterr().out


def test_cli_writes_xlsx(tmp_path):
    src = tmp_path / "in.csv"
    src.write_text(SAMPLE, encoding="utf-8")
    out = tmp_path / "out.xlsx"

    assert main([str(src), "-o", str(out)]) == 0

    sheet = load_workbook(out).active
    assert [c.value for c in sheet[1]] == ["customer_id", "e_mail", "signup_date"]
    assert sheet.max_row == 3


def test_cli_rejects_unknown_column(tmp_path, capsys):
    src = tmp_path / "in.csv"
    src.write_text(SAMPLE, encoding="utf-8")

    assert main([str(src), "--dates", "created"]) == 2
    assert "unknown column(s): created" in capsys.readouterr().err


def test_cli_reads_cp1252(tmp_path):
    src = tmp_path / "in.csv"
    src.write_bytes("name,city\nRen\xe9,Br\xfcnn\n".encode("cp1252"))
    out = tmp_path / "out.csv"

    assert main([str(src), "-o", str(out)]) == 0
    assert _read(out) == [{"name": "René", "city": "Brünn"}]


def test_sample_generator_is_deterministic(tmp_path):
    from scripts.make_sample import main as make

    a, b = tmp_path / "a.csv", tmp_path / "b.csv"
    make(a)
    make(b)
    assert a.read_bytes() == b.read_bytes()
